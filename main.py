import json
import io
from datetime import datetime
from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, send_file
from flask_login import login_required, current_user
from extensions import db
from models import Ingredient, Recipe, RecipeIngredient, Product, Sale, UNIDADES_POR_CATEGORIA
from currency import get_current_rate, refresh_rate

main_bp = Blueprint("main", __name__)
WHATSAPP_NUMERO = "584220143659"  # formato internacional para wa.me, sin '+' ni espacios


def _ingredientes_json(ingredientes):
    return json.dumps([
        {
            "id": i.id,
            "nombre": i.nombre,
            "categoria": i.categoria,
            "unidad_compra": i.unidad_compra,
            "costo_unitario_base": round(i.costo_unitario_base, 8),
        }
        for i in ingredientes
    ])


def _receta_items_json(receta):
    if not receta:
        return "[]"
    return json.dumps([
        {"ingredient_id": it.ingredient_id, "cantidad": it.cantidad_usada, "unidad": it.unidad_usada}
        for it in receta.items
    ])


# ---------- Tasa de cambio (API interna para el JS) ----------
@main_bp.route("/api/tasa")
@login_required
def api_tasa():
    rate = get_current_rate()
    if rate is None:
        return jsonify({"tasa": None, "fecha": None, "error": "No se pudo obtener la tasa BCV."})
    return jsonify({"tasa": rate.tasa_usd_ves, "fecha": rate.fecha,
                     "actualizado": rate.updated_at.isoformat()})


@main_bp.route("/api/tasa/actualizar", methods=["POST"])
@login_required
def api_tasa_actualizar():
    tasa = refresh_rate()
    if tasa is None:
        return jsonify({"ok": False, "error": "No se pudo consultar la fuente BCV en este momento."}), 502
    return jsonify({"ok": True, "tasa": tasa})


# ---------- Inventario ----------
@main_bp.route("/inventario")
@login_required
def inventario():
    ingredientes = Ingredient.query.filter_by(user_id=current_user.id).order_by(Ingredient.nombre).all()
    return render_template("inventory.html", ingredientes=ingredientes)


@main_bp.route("/inventario/nuevo", methods=["POST"])
@login_required
def inventario_nuevo():
    unidad = request.form.get("unidad_compra", "unidad").strip()
    ing = Ingredient(
        user_id=current_user.id,
        nombre=request.form.get("nombre", "").strip(),
        presentacion_cantidad=float(request.form.get("presentacion_cantidad") or 1),
        unidad_compra=unidad,
        presentacion_unidad=unidad,
        precio_compra_usd=float(request.form.get("precio_compra_usd") or 0),
        stock_actual=float(request.form.get("stock_actual") or 0),
        stock_minimo=float(request.form.get("stock_minimo") or 0),
    )
    db.session.add(ing)
    db.session.commit()
    flash(f"Ingrediente '{ing.nombre}' agregado.", "ok")
    return redirect(url_for("main.inventario"))


@main_bp.route("/inventario/<int:ing_id>/editar", methods=["POST"])
@login_required
def inventario_editar(ing_id):
    ing = Ingredient.query.filter_by(id=ing_id, user_id=current_user.id).first_or_404()
    ing.nombre = request.form.get("nombre", ing.nombre).strip()
    ing.presentacion_cantidad = float(request.form.get("presentacion_cantidad") or 1)
    unidad = request.form.get("unidad_compra", ing.unidad_compra).strip()
    ing.unidad_compra = unidad
    ing.presentacion_unidad = unidad
    ing.precio_compra_usd = float(request.form.get("precio_compra_usd") or 0)
    ing.stock_actual = float(request.form.get("stock_actual") or 0)
    ing.stock_minimo = float(request.form.get("stock_minimo") or 0)
    db.session.commit()
    flash(f"Ingrediente '{ing.nombre}' actualizado.", "ok")
    return redirect(url_for("main.inventario"))


@main_bp.route("/inventario/<int:ing_id>/eliminar", methods=["POST"])
@login_required
def inventario_eliminar(ing_id):
    ing = Ingredient.query.filter_by(id=ing_id, user_id=current_user.id).first_or_404()
    db.session.delete(ing)
    db.session.commit()
    flash("Ingrediente eliminado.", "ok")
    return redirect(url_for("main.inventario"))


# ---------- Calculadora / Recetas ----------
@main_bp.route("/")
@login_required
def calculadora():
    ingredientes = Ingredient.query.filter_by(user_id=current_user.id).order_by(Ingredient.nombre).all()
    recetas = Recipe.query.filter_by(user_id=current_user.id).order_by(Recipe.creado.desc()).all()
    return render_template("dashboard.html", ingredientes=ingredientes, recetas=recetas, receta=None,
                            ingredientes_json=_ingredientes_json(ingredientes),
                            receta_items_json="[]",
                            unidades_json=json.dumps(UNIDADES_POR_CATEGORIA))


@main_bp.route("/receta/<int:recipe_id>")
@login_required
def ver_receta(recipe_id):
    ingredientes = Ingredient.query.filter_by(user_id=current_user.id).order_by(Ingredient.nombre).all()
    recetas = Recipe.query.filter_by(user_id=current_user.id).order_by(Recipe.creado.desc()).all()
    receta = Recipe.query.filter_by(id=recipe_id, user_id=current_user.id).first_or_404()
    return render_template("dashboard.html", ingredientes=ingredientes, recetas=recetas, receta=receta,
                            ingredientes_json=_ingredientes_json(ingredientes),
                            receta_items_json=_receta_items_json(receta),
                            unidades_json=json.dumps(UNIDADES_POR_CATEGORIA))


@main_bp.route("/receta/guardar", methods=["POST"])
@login_required
def guardar_receta():
    recipe_id = request.form.get("recipe_id")
    if recipe_id:
        receta = Recipe.query.filter_by(id=recipe_id, user_id=current_user.id).first_or_404()
        receta.items.clear()
    else:
        receta = Recipe(user_id=current_user.id)
        db.session.add(receta)

    receta.nombre = request.form.get("nombre", "Receta sin nombre").strip()
    receta.porciones = float(request.form.get("porciones") or 1)
    receta.empaques_usd = float(request.form.get("empaques_usd") or 0)
    receta.otros_gastos_usd = float(request.form.get("otros_gastos_usd") or 0)
    receta.minutos_trabajo = float(request.form.get("minutos_trabajo") or 0)
    receta.valor_hora_usd = float(request.form.get("valor_hora_usd") or 0)
    receta.margen_pct = float(request.form.get("margen_pct") or 0)
    receta.comision_pct = float(request.form.get("comision_pct") or 0)

    ing_ids = request.form.getlist("ing_id")
    ing_cants = request.form.getlist("ing_cantidad")
    ing_unidades = request.form.getlist("ing_unidad")
    for ing_id, cantidad, unidad in zip(ing_ids, ing_cants, ing_unidades):
        if not ing_id or not cantidad:
            continue
        ingrediente = Ingredient.query.filter_by(id=int(ing_id), user_id=current_user.id).first()
        if ingrediente:
            receta.items.append(RecipeIngredient(
                ingredient=ingrediente, cantidad_usada=float(cantidad), unidad_usada=unidad or "unidad"
            ))

    db.session.commit()
    flash(f"Receta '{receta.nombre}' guardada.", "ok")
    return redirect(url_for("main.ver_receta", recipe_id=receta.id))


@main_bp.route("/receta/<int:recipe_id>/eliminar", methods=["POST"])
@login_required
def eliminar_receta(recipe_id):
    receta = Recipe.query.filter_by(id=recipe_id, user_id=current_user.id).first_or_404()
    db.session.delete(receta)
    db.session.commit()
    flash("Receta eliminada.", "ok")
    return redirect(url_for("main.calculadora"))


@main_bp.route("/receta/<int:recipe_id>/producir", methods=["POST"])
@login_required
def producir_receta(recipe_id):
    receta = Recipe.query.filter_by(id=recipe_id, user_id=current_user.id).first_or_404()
    lotes = float(request.form.get("lotes") or 1)
    insuficientes = receta.descontar_stock(lotes=lotes)
    db.session.commit()
    if insuficientes:
        flash(
            f"Producción registrada, pero se quedaron sin stock suficiente: {', '.join(insuficientes)}.",
            "error",
        )
    else:
        flash(f"Producción registrada: se descontó el inventario para {lotes} lote(s) de '{receta.nombre}'.", "ok")
    return redirect(url_for("main.ver_receta", recipe_id=receta.id))


# ---------- Ventas ----------
@main_bp.route("/ventas")
@login_required
def ventas():
    lista = Sale.query.filter_by(user_id=current_user.id).order_by(Sale.fecha.desc()).all()
    recetas = Recipe.query.filter_by(user_id=current_user.id).order_by(Recipe.nombre).all()
    total = sum(v.total_usd for v in lista)
    return render_template("sales.html", ventas=lista, recetas=recetas, total=total)


@main_bp.route("/ventas/nueva", methods=["POST"])
@login_required
def ventas_nueva():
    recipe_id = request.form.get("recipe_id") or None
    receta = None
    if recipe_id:
        receta = Recipe.query.filter_by(id=int(recipe_id), user_id=current_user.id).first()

    nombre_producto = request.form.get("producto_nombre", "").strip()
    if not nombre_producto and receta:
        nombre_producto = receta.nombre
    if not nombre_producto:
        nombre_producto = "Producto"

    venta = Sale(
        user_id=current_user.id,
        recipe_id=receta.id if receta else None,
        producto_nombre=nombre_producto,
        cantidad=float(request.form.get("cantidad") or 1),
        precio_unitario_usd=float(request.form.get("precio_unitario_usd") or 0),
        cliente=request.form.get("cliente", "").strip(),
        notas=request.form.get("notas", "").strip(),
    )
    db.session.add(venta)
    db.session.commit()
    flash("Venta registrada.", "ok")
    return redirect(url_for("main.ventas"))


@main_bp.route("/ventas/<int:sale_id>/eliminar", methods=["POST"])
@login_required
def ventas_eliminar(sale_id):
    venta = Sale.query.filter_by(id=sale_id, user_id=current_user.id).first_or_404()
    db.session.delete(venta)
    db.session.commit()
    flash("Venta eliminada.", "ok")
    return redirect(url_for("main.ventas"))


# ---------- Catálogo (productos con foto, cara al público) ----------
@main_bp.route("/productos")
@login_required
def productos():
    lista = Product.query.filter_by(user_id=current_user.id).order_by(Product.orden, Product.nombre).all()
    recetas = Recipe.query.filter_by(user_id=current_user.id).order_by(Recipe.nombre).all()
    return render_template("products.html", productos=lista, recetas=recetas)


@main_bp.route("/productos/nuevo", methods=["POST"])
@login_required
def productos_nuevo():
    recipe_id = request.form.get("recipe_id") or None
    prod = Product(
        user_id=current_user.id,
        recipe_id=int(recipe_id) if recipe_id else None,
        nombre=request.form.get("nombre", "").strip(),
        descripcion=request.form.get("descripcion", "").strip(),
        precio_usd=float(request.form.get("precio_usd") or 0),
        imagen_url=request.form.get("imagen_url", "").strip(),
        activo=bool(request.form.get("activo")),
    )
    db.session.add(prod)
    db.session.commit()
    flash(f"Producto '{prod.nombre}' agregado al catálogo.", "ok")
    return redirect(url_for("main.productos"))


@main_bp.route("/productos/<int:prod_id>/editar", methods=["POST"])
@login_required
def productos_editar(prod_id):
    prod = Product.query.filter_by(id=prod_id, user_id=current_user.id).first_or_404()
    recipe_id = request.form.get("recipe_id") or None
    prod.recipe_id = int(recipe_id) if recipe_id else None
    prod.nombre = request.form.get("nombre", prod.nombre).strip()
    prod.descripcion = request.form.get("descripcion", "").strip()
    prod.precio_usd = float(request.form.get("precio_usd") or 0)
    prod.imagen_url = request.form.get("imagen_url", "").strip()
    prod.activo = bool(request.form.get("activo"))
    db.session.commit()
    flash(f"Producto '{prod.nombre}' actualizado.", "ok")
    return redirect(url_for("main.productos"))


@main_bp.route("/productos/<int:prod_id>/eliminar", methods=["POST"])
@login_required
def productos_eliminar(prod_id):
    prod = Product.query.filter_by(id=prod_id, user_id=current_user.id).first_or_404()
    db.session.delete(prod)
    db.session.commit()
    flash("Producto eliminado del catálogo.", "ok")
    return redirect(url_for("main.productos"))


@main_bp.route("/catalogo")
def catalogo():
    # Público — sin login. Muestra el catálogo del único usuario de la app.
    from models import User
    user = User.query.first()
    productos = []
    if user:
        productos = Product.query.filter_by(user_id=user.id, activo=True).order_by(Product.orden, Product.nombre).all()
    return render_template("catalog.html", productos=productos, whatsapp=WHATSAPP_NUMERO)


# ---------- Reportes ----------
@main_bp.route("/reportes")
@login_required
def reportes():
    return render_template("reports.html")


@main_bp.route("/reportes/recetas.pdf")
@login_required
def reporte_recetas_pdf():
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm

    recetas = Recipe.query.filter_by(user_id=current_user.id).order_by(Recipe.nombre).all()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    elementos = [Paragraph("Veronik Sweet — Costos y precios por receta", styles["Title"]),
                 Paragraph(datetime.now().strftime("Generado el %d/%m/%Y"), styles["Normal"]),
                 Spacer(1, 0.6 * cm)]

    data = [["Receta", "Costo/porción", "Precio sugerido", "Ganancia/porción", "Margen"]]
    for r in recetas:
        precio = r.precio_sugerido_usd()
        ganancia = r.ganancia_por_porcion_usd()
        data.append([
            r.nombre,
            f"${r.costo_por_porcion_usd():.2f}",
            f"${precio:.2f}" if precio is not None else "—",
            f"${ganancia:.2f}" if ganancia is not None else "—",
            f"{r.margen_pct:.0f}%",
        ])

    tabla = Table(data, colWidths=[5.5 * cm, 3 * cm, 3.2 * cm, 3.2 * cm, 2 * cm])
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#c0316a")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e0b8c8")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fdf1f6")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    elementos.append(tabla)
    doc.build(elementos)
    buf.seek(0)
    return send_file(buf, mimetype="application/pdf", as_attachment=True,
                      download_name="veronik-sweet-costos.pdf")


@main_bp.route("/reportes/ventas.xlsx")
@login_required
def reporte_ventas_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    ventas_list = Sale.query.filter_by(user_id=current_user.id).order_by(Sale.fecha.desc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"

    headers = ["Fecha", "Producto", "Cliente", "Cantidad", "Precio unitario", "Total", "Notas"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C0316A")

    total_general = 0
    for v in ventas_list:
        ws.append([
            v.fecha.strftime("%Y-%m-%d %H:%M"),
            v.producto_nombre,
            v.cliente,
            v.cantidad,
            v.precio_unitario_usd,
            v.total_usd,
            v.notas,
        ])
        total_general += v.total_usd

    ws.append([])
    ws.append(["", "", "", "", "TOTAL", total_general, ""])
    ws["E" + str(ws.max_row)].font = Font(bold=True)
    ws["F" + str(ws.max_row)].font = Font(bold=True)

    for col, width in zip("ABCDEFG", [17, 24, 18, 10, 14, 12, 26]):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="veronik-sweet-ventas.xlsx",
    )
